"""Satu perintah untuk regenerate semuanya:  python build.py

Urutan: layar purwarupa -> gambar -> checks -> DOCX -> PDF (lewat Word) -> pratinjau PNG.
Opsi:  --skip-figures   lewati pembuatan ulang gambar (lebih cepat saat hanya mengedit teks)
       --no-pdf         hanya membuat DOCX
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import checks
from content import BODY, COVER, LAMPIRAN
from references import REFERENCES, REFERENCES_LAMPIRAN

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
OUT = ROOT / "output"
FIG = HERE / "figures"
PREVIEW = HERE / "_preview"
NAME = "Proposal_DAQ_Makhraj"

FONT = "Times New Roman"
BODY_PT = 12
TABLE_PT = 10
CAPTION_PT = 12
LINE_SPACING = 1.15
MARGINS = dict(left=2.5, right=2.5, top=2.5, bottom=2.5)  # cm
# Kolom angka/status yang diratakan tengah (kolom pertama selalu rata kiri).
CENTER_TABLES = {"kompetitor", "biaya", "skenario"}


# ------------------------------------------------------------------ helpers
def set_run_font(run, size=BODY_PT, bold=None, italic=None, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)
    szcs = OxmlElement("w:szCs")
    szcs.set(qn("w:val"), str(int(size * 2)))
    rpr.append(szcs)
    if bold is not None:
        run.bold = bold
        if bold:
            rpr.append(OxmlElement("w:bCs"))
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


TOKEN = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")


ARABIC_RUN = re.compile(r"([\u0600-\u06FF]+)")


def add_rich(par, text, size=BODY_PT, bold=False, italic=False):
    # Huruf Arab ditaruh di run tersendiri bertanda w:rtl. Word mengabaikan LRM, sehingga tanpa ini
    # daftar "ح-ه, ث-س" tercetak terbalik (diuji: run terpisah + w:rtl satu-satunya yang benar).
    for part in TOKEN.split(text):
        if not part:
            continue
        b, i = bold, italic
        if part.startswith("**") and part.endswith("**"):
            part, b = part[2:-2], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 1:
            part, i = part[1:-1], True
        for piece in ARABIC_RUN.split(part):
            if not piece:
                continue
            run = par.add_run(piece)
            set_run_font(run, size, bold=b, italic=i)
            if ARABIC_RUN.fullmatch(piece):
                run._element.get_or_add_rPr().append(OxmlElement("w:rtl"))


def fmt(par, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first=0.0, left=0.0, hanging=0.0,
        before=0, after=3, keep_next=False, spacing=LINE_SPACING):
    pf = par.paragraph_format
    pf.alignment = align
    pf.line_spacing = spacing
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.left_indent = Cm(left)
    if hanging:
        pf.first_line_indent = Cm(-hanging)
    else:
        pf.first_line_indent = Cm(first)
    pf.keep_with_next = keep_next
    pf.widow_control = True


def cell_margins(table, top=40, bottom=40, left=80, right=80):
    tblpr = table._element.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblpr.append(mar)


def no_split(row, header=False):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trpr.append(el)
    if header:
        trpr.append(OxmlElement("w:tblHeader"))


def page_number_field(par):
    run = par.add_run()
    set_run_font(run, 11)
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._element.append(el)


def set_page(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(MARGINS["left"])
    section.right_margin = Cm(MARGINS["right"])
    section.top_margin = Cm(MARGINS["top"])
    section.bottom_margin = Cm(MARGINS["bottom"])
    section.footer_distance = Cm(1.2)


# ------------------------------------------------------------------ numbering
def numbering():
    # Badan proposal bernomor 1, 2, ...; lampiran berhuruf A, B, ... (seperti contoh executive summary).
    figs, tabs = {}, {}
    for blocks, label in ((BODY, str), (LAMPIRAN, lambda n: chr(64 + n))):
        nf = nt = 0
        for b in blocks:
            if b[0] == "fig":
                nf += 1
                figs[b[1]] = label(nf)
            elif b[0] == "table":
                nt += 1
                tabs[b[1]] = label(nt)
    return figs, tabs


def resolve(text, figs, tabs):
    text = re.sub(r"\{fig:([a-z_]+)\}", lambda m: str(figs[m.group(1)]), text)
    return re.sub(r"\{tab:([a-z_]+)\}", lambda m: str(tabs[m.group(1)]), text)


# ------------------------------------------------------------------ builders
def build_cover(doc):
    def line(text, size, bold=False, before=0, after=0, color=None, italic=False):
        for i, t in enumerate(text.split("\n")):
            p = doc.add_paragraph()
            fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=before if i == 0 else 0, after=after, spacing=1.15)
            r = p.add_run(t)
            set_run_font(r, size, bold=bold, italic=italic, color=color)

    line(COVER["event"], 14, bold=True, before=36, after=0)
    p = doc.add_paragraph()
    fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=60, after=18)
    p.add_run().add_picture(str(FIG / "cover_icon.png"), width=Cm(4.2))
    line(COVER["title"], 26, bold=True, after=8, color="0F5E5E")
    line(COVER["subtitle"], 13, italic=False, after=0)
    line(COVER["authors_label"], 12, before=90, after=4)
    for a in COVER["authors"]:
        line(a, 12, bold=True)
    if COVER["participant"]:
        line(COVER["participant"], 12, before=8)
    line(COVER["footer"], 13, bold=True, before=120)


def build_body(doc, figs, tabs, blocks=BODY):
    for b in blocks:
        kind = b[0]
        if kind == "h1":
            p = doc.add_paragraph()
            fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=10, after=4, keep_next=True)
            add_rich(p, b[1].upper(), bold=True)
        elif kind == "h2":
            p = doc.add_paragraph()
            fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=6, after=2, keep_next=True)
            add_rich(p, b[1], bold=True)
        elif kind in ("p", "p0"):
            p = doc.add_paragraph()
            # Kalimat pengantar yang berakhir titik dua harus satu halaman dengan daftar berikutnya.
            fmt(p, first=1.0 if kind == "p" else 0.0, keep_next=b[1].rstrip().endswith((":", "berikut.")))
            add_rich(p, resolve(b[1], figs, tabs))
        elif kind == "list":
            items, style = b[1], b[2]
            for i, it in enumerate(items, 1):
                p = doc.add_paragraph()
                fmt(p, left=0.75, hanging=0.6, after=2)
                p.paragraph_format.tab_stops.add_tab_stop(Cm(0.75))
                mark = f"{i}." if style == "1" else "•"
                add_rich(p, f"{mark}\t" + resolve(it, figs, tabs))
            p.paragraph_format.space_after = Pt(5)
        elif kind == "fig":
            _, key, fname, caption, width = b
            p = doc.add_paragraph()
            fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=2, keep_next=True, spacing=1.0)
            p.add_run().add_picture(str(FIG / fname), width=Cm(width))
            c = doc.add_paragraph()
            fmt(c, align=WD_ALIGN_PARAGRAPH.CENTER, after=8, spacing=1.0)
            r = c.add_run(f"Gambar {figs[key]}. ")
            set_run_font(r, CAPTION_PT, bold=True)
            add_rich(c, caption, size=CAPTION_PT)
        elif kind == "table":
            _, key, caption, header, rows, widths = b
            c = doc.add_paragraph()
            fmt(c, align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=3, keep_next=True, spacing=1.0)
            r = c.add_run(f"Tabel {tabs[key]}. ")
            set_run_font(r, CAPTION_PT, bold=True)
            add_rich(c, caption, size=CAPTION_PT)
            t = doc.add_table(rows=1 + len(rows), cols=len(header))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            cell_margins(t)
            # Tabel ber-layout tetap memakai lebar tblGrid, bukan lebar sel; keduanya diisi.
            for ci, w in enumerate(widths):
                t.columns[ci].width = Cm(w)
            for ri, row in enumerate([header] + rows):
                no_split(t.rows[ri], header=(ri == 0))
                for ci, val in enumerate(row):
                    cell = t.cell(ri, ci)
                    cell.width = Cm(widths[ci])
                    cp = cell.paragraphs[0]
                    centered = ri > 0 and key in CENTER_TABLES and ci > 0
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    fmt(cp, align=WD_ALIGN_PARAGRAPH.CENTER if (ri == 0 or centered) else WD_ALIGN_PARAGRAPH.LEFT,
                        after=0, spacing=1.0, keep_next=ri < len(rows))
                    add_rich(cp, val, size=TABLE_PT, bold=(ri == 0))
            after = doc.add_paragraph()
            fmt(after, after=2, spacing=1.0)
            set_run_font(after.add_run(""), 6)
        elif kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build_lampiran(doc, figs, tabs):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    p = doc.add_paragraph()
    fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=6, keep_next=True)
    add_rich(p, "LAMPIRAN", bold=True)
    build_body(doc, figs, tabs, LAMPIRAN)
    # Sumber yang hanya dikutip di lampiran: dicetak di sini agar Daftar Pustaka utama tetap dalam batas halaman.
    if REFERENCES_LAMPIRAN:
        # Urut penulis pertama lalu tahun, supaya El Kheir 2025a, 2025b, 2026 berurutan.
        build_references(doc, REFERENCES_LAMPIRAN, "Rujukan Lampiran", key=lambda r: (r[2].split("(")[0].split(",")[0].lower(), r[1]))


def build_references(doc, refs=REFERENCES, title="DAFTAR PUSTAKA", key=lambda r: r[2].lower()):
    p = doc.add_paragraph()
    fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=12, after=6, keep_next=True)
    add_rich(p, title, bold=True)
    for _, _, text in sorted(refs, key=key):
        p = doc.add_paragraph()
        fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, left=1.0, hanging=1.0, after=2, spacing=1.0)
        add_rich(p, text)


def build_docx(path):
    figs, tabs = numbering()
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.element.rPr.rFonts.set(qn("w:cs"), FONT)

    cover = doc.sections[0]
    set_page(cover)
    build_cover(doc)

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    set_page(body)
    body.footer.is_linked_to_previous = False
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "1")
    body._sectPr.append(pg)
    fp = body.footer.paragraphs[0]
    fmt(fp, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    page_number_field(fp)

    build_body(doc, figs, tabs)
    build_references(doc)
    build_lampiran(doc, figs, tabs)
    doc.core_properties.title = "Makhraj: Proposal DAQ"
    doc.core_properties.author = ", ".join(COVER["authors"])
    doc.save(path)


def to_pdf(docx_path, pdf_path):
    # Instans Word TERPISAH (DispatchEx), supaya Word yang sedang dipakai pengguna tidak ikut ditutup.
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    try:
        word.Visible = False
        word.DisplayAlerts = 0
        doc = word.Documents.Open(str(docx_path), ReadOnly=True, AddToRecentFiles=False)
        try:
            doc.SaveAs(str(pdf_path), FileFormat=17)  # 17 = wdFormatPDF
        finally:
            doc.Close(False)
    finally:
        try:
            word.Quit()
        except (pythoncom.com_error, AttributeError):
            pass  # Word kadang sudah menutup diri setelah SaveAs; PDF sudah tersimpan.
        pythoncom.CoUninitialize()


def preview(pdf_path):
    import fitz

    PREVIEW.mkdir(exist_ok=True)
    for old in PREVIEW.glob("*.png"):
        old.unlink()
    doc = fitz.open(pdf_path)
    ref_page = lamp_page = None
    for i, page in enumerate(doc, 1):
        page.get_pixmap(dpi=80).save(PREVIEW / f"page{i:02d}.png")
        text = page.get_text()
        if ref_page is None and "DAFTAR PUSTAKA" in text:
            ref_page = i
        if lamp_page is None and "LAMPIRAN" in text:
            lamp_page = i
    total = len(doc)
    # Halaman bernomor = halaman PDF dikurangi sampul. Lampiran dimulai di halaman baru.
    print(f"  PDF: {total} halaman (1 sampul + {total - 1} halaman bernomor)")
    if ref_page and lamp_page:
        body_end, refs_end = ref_page - 1, lamp_page - 2
        print(f"  Bagian 1-10 berakhir di halaman bernomor {body_end} -> {'AMAN' if body_end <= 10 else 'LEBIH'}")
        print(f"  Bagian 1-10 + Daftar Pustaka berakhir di halaman bernomor {refs_end} -> "
              f"{'AMAN' if refs_end <= 10 else 'LEBIH'} (batas panduan 8-10 halaman)")
        print(f"  Lampiran: halaman bernomor {lamp_page - 1}-{total - 1}")
    return total


def main():
    args = set(sys.argv[1:])
    OUT.mkdir(exist_ok=True)
    if "--skip-figures" not in args:
        sys.path.insert(0, str(HERE / "screens"))
        import figures
        import render_screens

        render_screens.render()
        figures.build_all()
    checks.run()
    docx_path = OUT / f"{NAME}.docx"
    pdf_path = OUT / f"{NAME}.pdf"
    try:
        build_docx(docx_path)
    except PermissionError:
        # Berkas lama sedang terbuka (misalnya di Word): simpan sebagai salinan terbaru.
        docx_path = OUT / f"{NAME}-terbaru.docx"
        pdf_path = OUT / f"{NAME}-terbaru.pdf"
        print(f"  PERINGATAN: {NAME}.docx sedang terbuka; hasil disimpan sebagai {docx_path.name}")
        build_docx(docx_path)
    print(f"  DOCX: {docx_path}")
    if "--no-pdf" not in args:
        to_pdf(docx_path, pdf_path)
        print(f"  PDF : {pdf_path}")
        preview(pdf_path)


if __name__ == "__main__":
    main()
